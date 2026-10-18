#!/usr/bin/env python3
"""Real formation-top geometry benchmark; logs and final holdout are not scored."""
from __future__ import annotations
import argparse
from collections import Counter, defaultdict
from dataclasses import asdict
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import sys
import time

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'python/src'))
from earthforge import Engine, EngineError
from earthforge.horizons import Contact, fit_horizon_stack, predict_contacts, inside_hull, select_horizon_stack
import numpy as np
from openpyxl import load_workbook


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def save(path, value): Path(path).write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False)+'\n')


def assert_unchanged(snapshots):
    for path,expected in snapshots.items():
        if sha(path)!=expected:raise ValueError(f'Input or method changed during geometry study: {path}')


def install_access_guard():
    audit = {'logArraysOpened': 0, 'sourceLogArchivesOpened': 0}
    def guard(event, args):
        if event != 'open' or not isinstance(args[0], (str, bytes, os.PathLike)): return
        source=Path(os.fsdecode(args[0]))
        paths=[source,source.resolve()]
        if any(str(p).lower().endswith(('.npz', '.las')) or 'las_files_force_2020' in str(p).lower() or 'holdout' in {part.lower() for part in p.parts} for p in paths):
            raise PermissionError('Geometry validation must not open source logs or reserved arrays')
    sys.addaudithook(guard)
    return audit


def intake(book, sodir, plan):
    summary = json.loads((ROOT/'validation/force2020/intake-summary.json').read_text())
    split_path = ROOT/'validation/force2020/split.json'
    if sha(book) != summary['books']['groups']['sha256'] or sha(sodir) != summary['sodir']['sha256'] or sha(split_path) != summary['split_sha256']:
        raise ValueError('Input differs from frozen FORCE intake')
    split = json.loads(split_path.read_text())['wells']
    with sodir.open(encoding='utf-8-sig') as f:
        meta = {r['wlbWellboreName']: r for r in csv.DictReader(f)}
    names = plan['cohort']['tops_shallow_to_deep']
    picks = defaultdict(lambda: defaultdict(list)); skipped = 0
    w = load_workbook(book, read_only=True, data_only=True)
    rows = w['Sheet1'].iter_rows(values_only=True)
    if tuple(next(rows)[:8]) != ('Well identifier','Surface','X','Y','Z','MD','Interpreter','HorizonName'):
        raise ValueError('Unexpected group-top schema')
    for row in rows:
        name = str(row[0] or '')
        if split.get(name, {}).get('partition') != 'development':
            skipped += 1; continue
        if row[1] in names:
            picks[name][row[1]].append(row[:8])
    w.close()
    cohort, exclusions = [], []
    for well, member in sorted(split.items()):
        if member['partition'] != 'development': continue
        region = well.split('/')[0]
        reason = None; converted = []
        if region not in plan['cohort']['quadrants']: reason = 'outside-preselected-region'
        elif any(len(picks[well][h]) != 1 for h in names): reason = 'missing-or-ambiguous-group-top'
        else:
            try:
                for horizon in names:
                    row = picks[well][horizon][0]
                    x, y, z, md = map(float, row[2:6])
                    if not all(math.isfinite(v) for v in (x,y,z,md)) or not 0 <= -z < 10000 or md < 0: raise ValueError()
                    converted.append((Contact(well,member['group'],horizon,x,y,-z),md))
                if any(a[0].depth >= b[0].depth or a[1] >= b[1] for a,b in zip(converted,converted[1:])): reason = 'non-increasing-observed-stratigraphy'
                info = meta.get(well, {})
                water, rkb = float(info['wlbWaterDepth']), float(info['wlbKellyBushElevation'])
                first, md = converted[0]
                if water <= 0 or abs(first.depth-water) > 2 or abs(md-water-rkb) > 2: reason = 'unsupported-depth-reference'
            except (ValueError, TypeError, KeyError): reason = 'invalid-pick-or-no-exact-official-match'
        if reason: exclusions.append({'well':well,'reason':reason})
        else: cohort.extend(asdict(r) | {'region':region, 'block':well.split('-')[0]} for r,_ in converted)
    counts = {q: len({r['parent'] for r in cohort if r['region']==q}) for q in plan['cohort']['quadrants']}
    insufficient = {q for q,n in counts.items() if n < plan['cohort']['min_parents_per_region']}
    for well in sorted({r['well'] for r in cohort if r['region'] in insufficient}): exclusions.append({'well':well,'reason':'region-below-minimum-parent-count'})
    cohort = [r for r in cohort if r['region'] not in insufficient]
    # Coordinates, but no held-out depths, may define the common regional window.
    frames = {}
    for q in sorted({r['region'] for r in cohort}):
        xy = np.array([[r['x'],r['y']] for r in cohort if r['region']==q]); origin=xy.min(axis=0)-1
        extent=xy.max(axis=0)-origin+1
        frames[q]={'sourceXYOrigin':origin.tolist(),'bounds':{'min':dict(x=0.,y=0.,z=0.),'max':dict(x=float(extent[0]),y=float(extent[1]),z=10000.)}}
    for r in cohort:
        o=frames[r['region']]['sourceXYOrigin'];r['x']-=o[0];r['y']-=o[1]
    return cohort, {'inputs':{'groupsSha256':sha(book),'sodirSha256':sha(sodir),'splitSha256':sha(split_path)},'regions':frames,'parentsPerRegionBeforeMinimum':counts,'exclusions':exclusions,'holdoutWorkbookRowsSkippedBeforeCoordinateExtraction':skipped,'holdoutTargetsExtracted':0}


def contact(row): return Contact(**{k:row[k] for k in ('well','parent','horizon','x','y','depth')})


def folds(rows, design):
    key = 'parent' if design=='parent' else 'block'
    for group in sorted({r[key] for r in rows}):
        yield group, [r for r in rows if r[key]!=group], [r for r in rows if r[key]==group]


def evaluate(engine, cohort, audit, plan, out, methods):
    predictions=[]; ledger=[]; timings=[]
    names=plan['cohort']['tops_shallow_to_deep']
    for region in sorted(audit['regions']):
        rows=[r for r in cohort if r['region']==region]
        for design in ('parent','block'):
            for group, train, test in folds(rows,design):
                assert not ({r['parent'] for r in train}&{r['parent'] for r in test})
                enough=len({r['parent'] for r in train})>=plan['evaluation']['min_training_parents']
                targets=[r for r in test if names.index(r['horizon']) in plan['cohort']['primary_scored_units']]
                for method in methods:
                    start=time.perf_counter(); record={'region':region,'design':design,'heldGroup':group,'method':method,'trainingParents':sorted({r['parent'] for r in train}),'testParents':sorted({r['parent'] for r in test})}
                    error=None
                    if enough:
                        if method=='adaptive':
                            improvement=json.loads(Path(__file__).with_name('improvement-plan.json').read_text())
                            mapping={r['parent']:r['block'] for r in train} if design=='block' else None
                            history,fit=select_horizon_stack(engine,[contact(r) for r in train],names,audit['regions'][region]['bounds'],groups=mapping,score_horizons=[names[i] for i in plan['cohort']['primary_scored_units']],candidates=improvement['candidates_in_tie_preference_order'],min_training_parents=plan['evaluation']['min_training_parents'],missing_penalty_m=plan['evaluation']['missing_prediction_penalty_m'])
                        else:
                            history, fit=fit_horizon_stack([contact(r) for r in train],names,audit['regions'][region]['bounds'],method=method)
                        key=f'{region}-{design}-{group.replace("/","_").replace(" ","_")}-{method}'
                        history_file=out/'histories'/f'{key}.json';history_file.parent.mkdir(exist_ok=True)
                        save(history_file,history.to_record());record['history']=str(history_file.relative_to(out));record['modelSha256']=history.model_sha256;record['fit']=fit
                        try: got=predict_contacts(engine,history,[contact(r) for r in targets],names)
                        except EngineError as e:
                            error=str(e);got=[{'prediction':None,'status':'engine-error','occurrences':0} for r in targets]
                    else:
                        got=[{'prediction':None,'status':'insufficient-training-parents','occurrences':0} for r in targets]
                    elapsed=time.perf_counter()-start
                    record.update(elapsedSeconds=elapsed,engineError=error,scoredTargets=len(targets));ledger.append(record);timings.append(elapsed)
                    for target,pred in zip(targets,got):
                        local=[r for r in train if r['horizon']==target['horizon']]
                        xy=np.array([[r['x'],r['y']] for r in local]);point=np.array([target['x'],target['y']])
                        row={**target,**pred,'method':method,'design':design,'heldGroup':group,'insideTrainingHull':inside_hull(point,xy),'nearestTrainingDistanceM':float(np.linalg.norm(xy-point,axis=1).min())}
                        row['error']=None if pred['prediction'] is None else pred['prediction']-target['depth'];predictions.append(row)
    return predictions,ledger,{'foldRuns':len(timings),'wallSecondsSum':sum(timings),'medianFoldSeconds':float(np.median(timings)),'maxFoldSeconds':max(timings),'scope':'Observed sequential orchestration, fit, native process, serialization and history-file output. Not controlled repeated throughput.'}


def score(rows, penalty=10000):
    groups=defaultdict(list)
    for r in rows: groups[r['parent']].append(r)
    valid=[r for r in rows if r['error'] is not None]
    by_parent=[np.mean([abs(r['error']) if r['error'] is not None else penalty for r in v]) for v in groups.values()]
    conditional=defaultdict(list)
    for r in valid: conditional[r['parent']].append(r['error'])
    def metric(fn):return float(np.mean([np.mean(fn(np.array(v))) for v in conditional.values()])) if conditional else None
    return {'targets':len(rows),'parents':len(groups),'predicted':len(valid),'missing':len(rows)-len(valid),'coverage':len(valid)/len(rows) if rows else None,'penalizedMaeM':float(np.mean(by_parent)) if by_parent else None,'biasM':metric(lambda v:v),'maeM':metric(np.abs),'rmseM':math.sqrt(metric(np.square)) if conditional else None}


def summarize(rows, plan):
    results={}
    for design in sorted({r['design'] for r in rows}):
        for method in sorted({r['method'] for r in rows}):
            subset=[r for r in rows if r['design']==design and r['method']==method]
            groups={}
            for field in ('region','horizon','insideTrainingHull'):
                groups[field]={str(k):score([r for r in subset if r[field]==k]) for k in sorted({r[field] for r in subset})}
            results[f'{design}/{method}']={'overall':score(subset),'groups':groups}
    return results


def bootstrap(rows,plan):
    config=plan['evaluation']['bootstrap'];result={}
    for design in sorted({r['design'] for r in rows}):
        subset=[r for r in rows if r['design']==design];parents=sorted({r['parent'] for r in subset})
        rng=np.random.default_rng(config['seed']);indices=rng.integers(0,len(parents),size=(config['repeats'],len(parents)))
        losses={}
        for method in sorted({r['method'] for r in subset}):
            loss=[]
            for parent in parents:
                cases=[r for r in subset if r['method']==method and r['parent']==parent]
                loss.append(np.mean([abs(r['error']) if r['error'] is not None else plan['evaluation']['missing_prediction_penalty_m'] for r in cases]))
            losses[method]=np.array(loss)
            means=losses[method][indices].mean(axis=1)
            result[f'{design}/{method}']={'penalizedMaeM':float(np.mean(loss)),'intervalM':np.percentile(means,config['percentiles']).tolist()}
        for method in losses:
            if method=='plane':continue
            delta=losses[method]-losses['plane'];sample=delta[indices].mean(axis=1)
            result[f'{design}/{method}-minus-plane']={'differenceM':float(delta.mean()),'intervalM':np.percentile(sample,config['percentiles']).tolist()}
    return {'config':config,'estimates':result}


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--book',type=Path,required=True);parser.add_argument('--sodir',type=Path,required=True);parser.add_argument('--engine',type=Path,required=True);parser.add_argument('--out',type=Path,required=True);parser.add_argument('--adaptive',action='store_true')
    args=parser.parse_args();args.out.mkdir(parents=True,exist_ok=False)
    access=install_access_guard()
    methods=[Path(__file__),ROOT/'python/src/earthforge/horizons.py']
    inputs=[args.book,args.sodir,ROOT/'validation/force2020/split.json',ROOT/'validation/force2020/intake-summary.json',Path(__file__).with_name('plan.json')]
    if args.adaptive:inputs.append(Path(__file__).with_name('improvement-plan.json'))
    snapshots={p:sha(p) for p in [*methods,*inputs]}
    plan=json.loads((Path(__file__).with_name('plan.json')).read_text())
    cohort,audit=intake(args.book,args.sodir,plan);save(args.out/'cohort.json',cohort);save(args.out/'intake.json',audit)
    engine=Engine(args.engine,threads=1)
    predictions,ledger,timing=evaluate(engine,cohort,audit,plan,args.out,['constant','plane']+(['adaptive'] if args.adaptive else []))
    assert_unchanged(snapshots)
    save(args.out/'predictions.json',predictions);save(args.out/'folds.json',ledger)
    result={'schema':'earthforge-field-geometry-results-v1','planSha256':snapshots[Path(__file__).with_name('plan.json')],'improvementPlanSha256':snapshots[Path(__file__).with_name('improvement-plan.json')] if args.adaptive else None,'engineSha256':engine.sha256,'sourceFiles':{str(p.relative_to(ROOT)):snapshots[p] for p in methods},'inputsAndMethodsUnchangedThroughRun':True,'intake':audit,'cohort':{'wells':len({r['well'] for r in cohort}),'parents':len({r['parent'] for r in cohort}),'regions':len(audit['regions']),'contactPicksIncludingSeafloor':len(cohort)},'results':summarize(predictions,plan),'bootstrap':bootstrap(predictions,plan),'timing':timing,'access':access,'scope':plan['scope_limits']}
    save(args.out/'results.json',result)
    for name,s in result['results'].items(): print(name,json.dumps(s['overall']),flush=True)
    print('Saved all predictions, fold histories, failures and input fingerprints.',flush=True)


if __name__=='__main__':main()
